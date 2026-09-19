"""
Admin Siswa Endpoints — CRUD & Excel Bulk Import
===============================================
Menyediakan REST API pengelolaan profil siswa untuk admin sekolah.
"""
import secrets
import structlog
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query
from pydantic import BaseModel, Field
from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.core.dependencies import get_current_admin
from app.models.admin import Admin
from app.models.siswa import Siswa
from app.schemas.sekolah_schemas import SiswaResponse, SiswaCreate, SiswaUpdate, SiswaList

logger = structlog.get_logger(__name__)
router = APIRouter()


# ─── ENDPOINTS SISWA ADMIN ────────────────────────────────────────────────────

@router.get("", response_model=SiswaList, tags=["Admin Siswa"])
async def get_students(
    sekolah_id: Optional[int] = None,
    kelas: Optional[str] = None,
    search: Optional[str] = None,
    page: int = Query(1, ge=1, description="Nomor halaman"),
    limit: int = Query(10, ge=1, le=2000, description="Jumlah data per halaman"),
    current_admin: Admin = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db)
):
    """
    Mengambil daftar seluruh siswa sekolah dengan filter opsional dan paginasi.
    """
    logger.info("👨‍🎓 Admin fetching students list", admin_id=current_admin.id, page=page, limit=limit)
    
    from sqlalchemy import or_
    query = select(Siswa)
    count_query = select(func.count()).select_from(Siswa)
    
    # Force tenant isolation for non-super_admin
    if current_admin.role != "super_admin":
        sekolah_id = current_admin.sekolah_id
        
    if sekolah_id:
        query = query.where(Siswa.sekolah_id == sekolah_id)
        count_query = count_query.where(Siswa.sekolah_id == sekolah_id)
    if kelas:
        query = query.where(Siswa.kelas == kelas)
        count_query = count_query.where(Siswa.kelas == kelas)
    if search:
        search_term = f"%{search}%"
        search_filter = or_(
            Siswa.nama_lengkap.ilike(search_term),
            Siswa.nis.ilike(search_term),
            Siswa.nisn.ilike(search_term)
        )
        query = query.where(search_filter)
        count_query = count_query.where(search_filter)
        
    # Get total
    total_result = await db.execute(count_query)
    total = total_result.scalar() or 0
    
    # Pagination
    query = query.offset((page - 1) * limit).limit(limit)
    
    result = await db.execute(query)
    students = result.scalars().all()
    
    return {
        "items": students,
        "total": total,
        "page": page,
        "size": limit
    }


@router.get("/kelas", response_model=list[str], tags=["Admin Siswa"])
async def get_classes(
    sekolah_id: Optional[int] = None,
    current_admin: Admin = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db)
):
    """
    Mengambil daftar kelas unik yang ada pada data siswa.
    """
    query = select(Siswa.kelas).distinct().where(Siswa.kelas.is_not(None))
    # Force tenant isolation for non-super_admin
    if current_admin.role != "super_admin":
        sekolah_id = current_admin.sekolah_id

    if sekolah_id:
        query = query.where(Siswa.sekolah_id == sekolah_id)
        
    result = await db.execute(query)
    classes = result.scalars().all()
    # Hapus string kosong dan urutkan
    return sorted([c for c in classes if c.strip()])


class BulkKelulusanRequest(BaseModel):
    angkatan: int = Field(..., ge=1900, le=2100, description="Tahun angkatan siswa")
    tahun_lulus: Optional[int] = Field(None, ge=1900, le=2100, description="Tahun lulus (null untuk reset ke Siswa Aktif)")
    kelas: Optional[str] = Field(None, description="Filter kelas opsional")


@router.post("/kelulusan-massal", tags=["Admin Siswa"])
async def bulk_update_kelulusan(
    payload: BulkKelulusanRequest,
    current_admin: Admin = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db)
):
    """
    Mengubah status kelulusan (tahun_lulus) secara massal berdasarkan angkatan (dan kelas opsional).
    """
    logger.info("👨‍🎓 Admin updating bulk kelulusan", admin_id=current_admin.id, angkatan=payload.angkatan)
    
    query = select(Siswa).where(Siswa.angkatan == payload.angkatan)
    
    if current_admin.role != "super_admin":
        query = query.where(Siswa.sekolah_id == current_admin.sekolah_id)
        
    if payload.kelas and payload.kelas.strip():
        query = query.where(Siswa.kelas == payload.kelas.strip())
        
    result = await db.execute(query)
    students = result.scalars().all()
    
    if not students:
        raise HTTPException(
            status_code=404, 
            detail=f"Tidak ada data siswa ditemukan untuk angkatan {payload.angkatan}" + (f" kelas {payload.kelas}" if payload.kelas else "")
        )
        
    updated_count = 0
    from datetime import datetime, timezone
    now = datetime.now(timezone.utc)
    
    for s in students:
        s.tahun_lulus = payload.tahun_lulus
        s.updated_at = now
        updated_count += 1
        
    await db.commit()
    
    status_str = f"LULUS ({payload.tahun_lulus})" if payload.tahun_lulus else "SISWA AKTIF"
    logger.info("✅ Bulk kelulusan completed", updated_count=updated_count, status_str=status_str)
    
    return {
        "status": "success",
        "message": f"Berhasil memperbarui status {updated_count} siswa angkatan {payload.angkatan} menjadi {status_str}!",
        "updated_count": updated_count,
        "angkatan": payload.angkatan,
        "tahun_lulus": payload.tahun_lulus
    }


@router.post("", response_model=SiswaResponse, status_code=201, tags=["Admin Siswa"])
async def create_student(
    payload: SiswaCreate,
    current_admin: Admin = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db)
):
    """
    Menambahkan siswa baru ke sistem.
    Secara otomatis menghasilkan `entropy_seed` acak kriptografis sepanjang 32-karakter hex (16 bytes)
    untuk kebutuhan modul NeuralKeyGen.
    """
    logger.info("👨‍🎓 Admin creating student", admin_id=current_admin.id, nis=payload.nis)
    
    # 1. Cek duplikasi NIS
    query_nis = select(Siswa).where(Siswa.nis == payload.nis)
    res_nis = await db.execute(query_nis)
    if res_nis.scalar_one_or_none():
        raise HTTPException(
            status_code=400,
            detail=f"Siswa dengan NIS {payload.nis} sudah terdaftar di sistem"
        )
        
    # 2. Hasilkan entropy_seed acak (16 bytes = 32 hex chars)
    entropy_seed = secrets.token_hex(16)
    
    final_sekolah_id = payload.sekolah_id
    if current_admin.role != "super_admin":
        final_sekolah_id = current_admin.sekolah_id

    # 3. Simpan ke database
    new_student = Siswa(
        nis=payload.nis,
        nisn=payload.nisn,
        nama_lengkap=payload.nama_lengkap,
        tgl_lahir=payload.tgl_lahir,
        tempat_lahir=payload.tempat_lahir,
        jenis_kelamin=payload.jenis_kelamin,
        agama=payload.agama,
        alamat=payload.alamat,
        kelas=payload.kelas,
        jurusan=payload.jurusan,
        angkatan=payload.angkatan,
        tahun_lulus=payload.tahun_lulus,
        email=payload.email,
        telepon=payload.telepon,
        telepon_ortu=payload.telepon_ortu,
        nama_ortu=payload.nama_ortu,
        sekolah_id=final_sekolah_id,
        entropy_seed=entropy_seed,
        is_active=True
    )
    
    db.add(new_student)
    await db.commit()
    await db.refresh(new_student)
    
    logger.info("✅ Student created successfully", siswa_id=new_student.id, nis=new_student.nis)
    return new_student


@router.get("/{id}", response_model=SiswaResponse, tags=["Admin Siswa"])
async def get_student_detail(
    id: int,
    current_admin: Admin = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db)
):
    """
    Mengambil informasi profil lengkap siswa berdasarkan ID.
    """
    logger.info("👨‍🎓 Admin fetching student detail", admin_id=current_admin.id, siswa_id=id)
    
    query = select(Siswa).where(Siswa.id == id)
    if current_admin.role != "super_admin":
        query = query.where(Siswa.sekolah_id == current_admin.sekolah_id)
        
    result = await db.execute(query)
    siswa = result.scalar_one_or_none()
    
    if not siswa:
        raise HTTPException(status_code=404, detail="Siswa tidak ditemukan atau di luar wewenang sekolah Anda")
        
    return siswa


@router.put("/{id}", response_model=SiswaResponse, tags=["Admin Siswa"])
async def update_student(
    id: int,
    payload: SiswaUpdate,
    current_admin: Admin = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db)
):
    """
    Memperbarui profil informasi siswa.
    """
    logger.info("👨‍🎓 Admin updating student", admin_id=current_admin.id, siswa_id=id)
    
    query = select(Siswa).where(Siswa.id == id)
    if current_admin.role != "super_admin":
        query = query.where(Siswa.sekolah_id == current_admin.sekolah_id)
        
    result = await db.execute(query)
    siswa = result.scalar_one_or_none()
    
    if not siswa:
        raise HTTPException(status_code=404, detail="Siswa tidak ditemukan atau di luar wewenang sekolah Anda")
        
    # Update fields
    update_data = payload.model_dump(exclude_unset=True)
    for key, value in update_data.items():
        setattr(siswa, key, value)
        
    await db.commit()
    await db.refresh(siswa)
    
    logger.info("✅ Student updated successfully", siswa_id=id)
    return siswa


@router.delete("/{id}", tags=["Admin Siswa"])
async def delete_student(
    id: int,
    current_admin: Admin = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db)
):
    """
    Menghapus siswa dari sistem (soft-delete / non-aktifkan akun siswa).
    """
    logger.info("👨‍🎓 Admin deleting student", admin_id=current_admin.id, siswa_id=id)
    
    query = select(Siswa).where(Siswa.id == id)
    if current_admin.role != "super_admin":
        query = query.where(Siswa.sekolah_id == current_admin.sekolah_id)
        
    result = await db.execute(query)
    siswa = result.scalar_one_or_none()
    
    if not siswa:
        raise HTTPException(status_code=404, detail="Siswa tidak ditemukan atau di luar wewenang sekolah Anda")
        
    siswa.is_active = False
    await db.commit()
    
    logger.info("✅ Student deactivated successfully", siswa_id=id)
    return {"status": "success", "message": f"Akun Siswa ID {id} berhasil dinonaktifkan"}


@router.post("/import", tags=["Admin Siswa"])
async def import_students_via_excel(
    file: UploadFile = File(..., description="File Excel (.xlsx/.xls) daftar siswa baru"),
    current_admin: Admin = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db)
):
    """
    Melakukan import massal (bulk import) data siswa baru via Excel.
    Secara otomatis menghasilkan `entropy_seed` unik untuk setiap baris siswa baru.

    Robust terhadap data kotor:
    - Placeholder seperti '-', '--', 'N/A', 'null' dianggap kosong (NULL).
    - Email divalidasi format; invalid/jika duplikat → diabaikan (NULL), baris tetap masuk.
    - Duplikat NIS/NISN/email (di DB maupun antar-baris) dilewati tanpa membatalkan import lain.
    - Mengembalikan laporan per-baris untuk data yang gagal.
    """
    logger.info("📥 Admin importing students via Excel", admin_id=current_admin.id, filename=file.filename)

    # Validasi extension file
    if not file.filename or not (file.filename.endswith('.xlsx') or file.filename.endswith('.xls')):
        raise HTTPException(status_code=400, detail="Hanya mendukung berkas Excel (.xlsx atau .xls)")

    import openpyxl
    import io
    import re
    from datetime import datetime
    from sqlalchemy.exc import IntegrityError, DataError

    contents = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Gagal membaca file Excel. Pastikan format .xlsx (Excel modern). "
                   "Jika file berformat .xls lama, simpan ulang sebagai .xlsx terlebih dahulu."
        )

    ws = wb.active
    headers = [str(cell.value).strip().lower() if cell.value is not None else "" for cell in ws[1]]

    if not headers or "nis" not in headers or "nama_lengkap" not in headers:
        raise HTTPException(
            status_code=400,
            detail="Format Excel tidak valid. Kolom wajib: 'nis' dan 'nama_lengkap'. "
                   "Pastikan menggunakan template yang disediakan."
        )

    EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
    # Nilai placeholder yang umum di Excel sekolah → dianggap kosong
    PLACEHOLDERS = {"-", "--", "---", "n/a", "na", "none", "null", "nil", "kosong", "belum ada", "x"}

    def clean_val(val, maxlen: Optional[int] = None) -> Optional[str]:
        """Normalisasi nilai sel: placeholder → None, float bulat → int string, opsional clip panjang."""
        if val is None:
            return None
        if isinstance(val, float) and val.is_integer():
            val = int(val)
        s = str(val).strip()
        if not s or s.lower() in PLACEHOLDERS:
            return None
        if maxlen and len(s) > maxlen:
            s = s[:maxlen]
        return s

    def clean_gender(val) -> Optional[str]:
        """Normalisasi jenis kelamin ke 'L'/'P' (kolom String(1))."""
        s = clean_val(val)
        if not s:
            return None
        s = s.upper()
        if s.startswith("L"):
            return "L"
        if s.startswith("P"):
            return "P"
        return None

    def clean_email(val) -> Optional[str]:
        s = clean_val(val)
        if not s:
            return None
        s = s.lower()
        return s if EMAIL_RE.match(s) else None

    def clean_phone(val) -> Optional[str]:
        s = clean_val(val)
        if not s:
            return None
        # Buang spasi/tanda hubung; sisakan digit dan prefix '+'
        s = re.sub(r"[\s\-()]", "", s)
        return s or None

    def parse_date(val):
        if isinstance(val, datetime):
            return val.date()
        from datetime import date as _date
        if isinstance(val, _date):
            return val
        if isinstance(val, str):
            for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
                try:
                    return datetime.strptime(val.strip(), fmt).date()
                except ValueError:
                    continue
        return None

    def to_int(val):
        s = clean_val(val)
        if not s:
            return None
        try:
            return int(float(s))
        except (ValueError, TypeError):
            return None

    target_sekolah_id = current_admin.sekolah_id or 1

    # Pre-fetch semua NIS/NISN/email sekolah ini + global unik untuk deteksi duplikat cepat
    res_existing = await db.execute(
        select(Siswa.nis, Siswa.nisn, Siswa.email).where(Siswa.sekolah_id == target_sekolah_id)
    )
    existing_nis = set()
    existing_nisn = set()
    existing_email = set()
    for nis_, nisn_, email_ in res_existing.all():
        if nis_: existing_nis.add(nis_)
        if nisn_: existing_nisn.add(nisn_)
        if email_: existing_email.add(email_.lower())

    res_global = await db.execute(select(Siswa.nisn, Siswa.email).where(
        (Siswa.nisn.is_not(None)) | (Siswa.email.is_not(None))
    ))
    for nisn_, email_ in res_global.all():
        if nisn_: existing_nisn.add(nisn_)
        if email_: existing_email.add(email_.lower())

    imported_count = 0
    skipped_rows: list[dict] = []
    total_processed = 0
    row_num = 1  # header adalah baris 1

    for row in ws.iter_rows(min_row=2, values_only=True):
        row_num += 1
        if not any(cell is not None and str(cell).strip() for cell in row):
            continue
        total_processed += 1

        row_dict = dict(zip(headers, row))
        nis = clean_val(row_dict.get("nis"), maxlen=20)
        if not nis:
            skipped_rows.append({"baris": row_num, "nis": None, "alasan": "Kolom NIS kosong"})
            continue

        nama = clean_val(row_dict.get("nama_lengkap"), maxlen=255) or "Tanpa Nama"

        # ── Deteksi duplikat (DB + antar-baris dalam file) ──────────────────
        if nis in existing_nis:
            skipped_rows.append({"baris": row_num, "nis": nis, "alasan": f"NIS {nis} sudah terdaftar"})
            continue

        nisn = clean_val(row_dict.get("nisn"))
        email = clean_email(row_dict.get("email"))
        if nisn and len(nisn) != 10:
            nisn = None  # NISN wajib 10 digit; selain itu dianggap data kotor (jangan dipotong)
        if nisn and nisn in existing_nisn:
            skipped_rows.append({"baris": row_num, "nis": nis, "alasan": f"NISN {nisn} sudah dipakai siswa lain"})
            continue
        if email and email in existing_email:
            email = None  # email duplikat tidak menggagalkan baris, cukup dikosongkan

        kelas = clean_val(row_dict.get("kelas"), maxlen=20)
        angkatan = to_int(row_dict.get("angkatan"))

        new_student = Siswa(
            nis=nis,
            nisn=nisn,
            nama_lengkap=nama,
            kelas=kelas,
            angkatan=angkatan,
            jurusan=clean_val(row_dict.get("jurusan"), maxlen=100),
            tgl_lahir=parse_date(row_dict.get("tgl_lahir")),
            tempat_lahir=clean_val(row_dict.get("tempat_lahir"), maxlen=100),
            jenis_kelamin=clean_gender(row_dict.get("jenis_kelamin")),
            agama=clean_val(row_dict.get("agama"), maxlen=20),
            email=email,
            telepon=clean_phone(row_dict.get("telepon")),
            nama_ortu=clean_val(row_dict.get("nama_ortu"), maxlen=255),
            telepon_ortu=clean_phone(row_dict.get("telepon_ortu")),
            alamat=clean_val(row_dict.get("alamat")),
            sekolah_id=target_sekolah_id,
            entropy_seed=secrets.token_hex(16),
            is_active=True
        )

        # Flush per-baris dengan savepoint agar satu baris rusak tidak menggagalkan semuanya
        try:
            async with db.begin_nested():
                db.add(new_student)
                await db.flush()
            imported_count += 1
            existing_nis.add(nis)
            if nisn: existing_nisn.add(nisn)
            if email: existing_email.add(email)
        except (IntegrityError, DataError) as e:
            # Savepoint sudah otomatis di-rollback oleh context manager.
            # Baris ini dilewati, baris lain tetap diproses.
            try:
                db.expunge(new_student)
            except Exception:
                pass
            constraint = getattr(getattr(e.orig, "diag", None), "constraint_name", None) or "data tidak valid"
            skipped_rows.append({
                "baris": row_num,
                "nis": nis,
                "alasan": f"Gagal simpan ({constraint})"
            })
            continue

    await db.commit()
    logger.info(
        "✅ Bulk import completed",
        total_imported=imported_count,
        total_skipped=len(skipped_rows),
        total_processed=total_processed
    )

    result = {
        "status": "success" if imported_count > 0 else "warning",
        "message": f"Berhasil mengimpor {imported_count} siswa baru. Dilewati: {len(skipped_rows)} baris.",
        "details": {
            "total_processed": total_processed,
            "total_imported": imported_count,
            "total_skipped": len(skipped_rows),
        }
    }
    if skipped_rows:
        result["details"]["errors"] = skipped_rows[:50]  # batasi 50 error pertama
    return result
